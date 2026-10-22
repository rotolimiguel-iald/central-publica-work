"""Bounded, offline artifact generation and authenticated-owner storage.

Contract:
    ArtifactStore(app_root) stores below app_root/.iald_runtime/task_artifacts.
    create(owner, kind, name, payload) returns {"status": "generated", "artifacts": [...]}.
    ingest(owner, filename, raw_bytes) returns a SINGLE artifact dictionary.
    Both include id, filename, sha256, bytes, url, mobile_url, expires_at (UTC).
    resolve(owner, id) returns (Path, metadata), including mime and download_headers.
    issue_ticket(owner, id) returns a mobile URL with a fixed 900-second lifetime.
    resolve_ticket(id, token) validates that capability without accepting an owner.
    Reissuing a ticket revokes its predecessor. Expiry is checked when resolved;
    it does not revoke bytes already downloaded or an already started response.

The app MUST supply owner from its authenticated session, never a tool argument,
and authorize delegation before calling this helper. This module has no HTTP
handler. A future handler must call resolve for every download, use the returned
    attachment headers (also for resolve_ticket), and never mount this directory
    as static content. Tokens are ephemeral bearer capabilities, so do not log
    mobile URLs or send them as referrers. No script,
formula, browser, renderer, shell, or network client is executed here. Formula
caches are not calculated and no visual QA or script test result is asserted.

Use an app-private root with OS access limited to the app account. Link/reparse
points and hardlinks are rejected. Hashes detect corruption, not an administrator
rewriting both bytes and metadata. A returned Path is not an open file capability:
the caller must not permit concurrent untrusted filesystem mutation before send.
Writes use an exclusive per-owner lock and an atomic directory rename. A crashed
writer may leave a lock/pending directory; later writes fail closed for manual
recovery, without guessing whether an uncertain write should be repeated.

Dependencies: Python >=3.10; openpyxl for XLSX; python-pptx and its lxml dependency
for PPTX; Pillow for PNG ingestion; pypdf for simple static PDF ingestion. Office
ingestion accepts only the documented, bounded, data-only subset (no embedded
assets). PDFs with object/xref streams, encryption or active features are refused.
Imports do not create storage. This is validation, not a malware scanner.
"""

from __future__ import annotations

import csv
import hashlib
import hmac
import html
import io
import json
import math
import os
import posixpath
import re
import secrets
import stat
import struct
import sys
import time
import unicodedata
import uuid
import zipfile
import zlib
from contextlib import contextmanager
from datetime import datetime, timezone
from functools import wraps
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree as ET


MAX_ARTIFACT_BYTES = 20 * 1024 * 1024
MAX_OWNER_ARTIFACTS = 5000
MAX_OWNER_BYTES = 2 * 1024 * 1024 * 1024
MAX_SHEETS = 16
MAX_ROWS = 10_000  # Excludes the header.
MAX_COLUMNS = 100
MAX_CELLS = 100_000  # Across the whole workbook or CSV.
MAX_SLIDES = 50
MAX_ZIP_ENTRIES = 2048
MAX_ZIP_MEMBER_BYTES = 16 * 1024 * 1024
MAX_ZIP_EXPANDED_BYTES = 60 * 1024 * 1024
MAX_ZIP_RATIO = 1000
MAX_METADATA_BYTES = 4096
EXTENSIONS = {"python": ".py", "text": ".txt", "json": ".json", "csv": ".csv",
              "spreadsheet": ".xlsx", "slides": ".pptx", "website": ".html"}
INGEST_EXTENSIONS = {**EXTENSIONS, "markdown": ".md", "png": ".png", "pdf": ".pdf"}
MIME_TYPES = {"python": "text/x-python", "text": "text/plain", "json": "application/json",
              "csv": "text/csv", "markdown": "text/markdown", "website": "text/html",
              "spreadsheet": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
              "slides": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
              "png": "image/png", "pdf": "application/pdf"}
TICKET_TTL_SECONDS = 900
PUBLIC_BASE_URL = "https://demo.iald.ia.br"


def _artifact_url(artifact_id):
    return PUBLIC_BASE_URL + "/api/task-artifacts/" + artifact_id

TOOL_DEFINITION = {
    "type": "function",
    "function": {
        "name": "create_artifact",
        "description": (
            "Generate one downloadable file, never run or publish it. Shapes: "
            "python/text {content:string}; json {data:JSON value}; "
            "csv {headers:[string],rows:[[scalar]]}; spreadsheet "
            "{sheets:[{name,headers:[string],rows:[[scalar]],formulas?:{B4:'=SUM(B2:B3)'}}]}; "
            "slides {slides:[{title:string,bullets:[string]}]}; website "
            "{title:string,sections:[{heading:string,text:string,bullets?:[string]}]} "
            "or {html:string} (restricted static HTML/CSS; no scripts, links or assets). "
            "Spreadsheet text stays text; only formulas maps create formulas. Allowed: "
            "SUM/AVERAGE/MIN/MAX and arithmetic, same-sheet references within the table. "
            "Formula targets must be EMPTY reserved data cells; headers are row 1 and data start at row 2. Never replace input values with formulas. No recalculation or rendering is performed. "
            "20 MiB/file; 5000 files and 2 GiB/owner; 16 sheets, 100 columns, 10000 data "
            "rows, 100000 total cells; 50 slides, 6 bullets/slide. Authenticated owner is supplied by the app. "
            "Show the authenticated url AND temporary mobile_url with its expires_at UTC timestamp. "
            "Mobile links expire after 15 minutes; report generated, never tests passed."
        ),
        "parameters": {
            "type": "object", "additionalProperties": False,
            "required": ["kind", "name", "payload"],
            "properties": {
                "kind": {"type": "string", "enum": list(EXTENSIONS)},
                "name": {"type": "string", "minLength": 1, "maxLength": 255,
                         "description": "Download name only; no path or separators."},
                "payload": {"type": "object", "description": "Use exactly the shape for kind described above."},
            },
        },
    },
}
CREATE_ARTIFACT_TOOL = TOOL_DEFINITION  # Backward-compatible descriptive alias.


class ArtifactError(ValueError):
    """Safe Portuguese error for presentation to an authenticated caller."""


def _fail(message):
    raise ArtifactError(message)


def _public(method):
    @wraps(method)
    def guarded(*args, **kwargs):
        try:
            return method(*args, **kwargs)
        except ArtifactError:
            raise
        except ImportError as exc:
            raise ArtifactError("Dependência indisponível: verifique openpyxl, python-pptx, Pillow e pypdf para o formato solicitado.") from exc
        except (OSError, ValueError, TypeError, KeyError, OverflowError, RecursionError,
                zipfile.BadZipFile, ET.ParseError, csv.Error) as exc:
            raise ArtifactError("Não foi possível validar ou armazenar o artefato local.") from exc
    return guarded


def _text(value, limit=MAX_ARTIFACT_BYTES):
    if type(value) is not str or len(value) > limit:
        _fail("Texto inválido ou acima do limite permitido.")
    if any((ord(c) < 32 and c not in "\t\n\r") or ord(c) == 127
           or 0xD800 <= ord(c) <= 0xDFFF or ord(c) in (0xFFFE, 0xFFFF) for c in value):
        _fail("O texto contém caracteres de controle inválidos.")
    return value


def _object(value, required, optional=()):
    if type(value) is not dict or not set(required) <= value.keys() or value.keys() - set(required) - set(optional):
        _fail("Estrutura inválida: confira os campos exigidos para o formato.")
    return value


def _json_tree(value):
    """Bound JSON depth, size and node count without invoking custom objects."""
    remaining = [MAX_ARTIFACT_BYTES, 200_000]
    active = set()

    def visit(item, depth):
        remaining[1] -= 1
        if depth > 32 or remaining[1] < 0:
            _fail("Conteúdo muito complexo para um único artefato.")
        if type(item) is str:
            remaining[0] -= len(_text(item).encode("utf-8"))
        elif item is None or type(item) is bool:
            remaining[0] -= 5
        elif type(item) in (int, float):
            if (type(item) is float and not math.isfinite(item)) or (type(item) is int and item.bit_length() > 1024):
                _fail("Número inválido ou fora do limite permitido.")
            remaining[0] -= len(str(item))
        elif type(item) in (dict, list):
            if id(item) in active:
                _fail("O conteúdo não pode conter referências circulares.")
            active.add(id(item))
            if type(item) is dict:
                for key, val in item.items():
                    if type(key) is not str:
                        _fail("As chaves do conteúdo devem ser texto.")
                    visit(key, depth + 1)
                    visit(val, depth + 1)
            else:
                for val in item:
                    visit(val, depth + 1)
            active.remove(id(item))
        else:
            _fail("O conteúdo deve conter apenas dados JSON, sem objetos executáveis.")
        if remaining[0] < 0:
            _fail("O conteúdo excede 20 MiB.")
    visit(value, 0)


def _load_json(raw):
    def pairs(items):
        result = {}
        for key, val in items:
            if key in result:
                _fail("JSON inválido: há chaves duplicadas.")
            result[key] = val
        return result

    def constant(_):
        _fail("JSON inválido: números não finitos não são permitidos.")
    result = json.loads(raw, object_pairs_hook=pairs, parse_constant=constant)
    _json_tree(result)
    return result


def _filename(name, kind):
    _text(name, 255)
    if not name.strip() or any(c in name for c in "/\\:") or any(ord(c) < 32 for c in name) or name in (".", ".."):
        _fail("Informe somente o nome do arquivo, sem caminho.")
    stem = name.rsplit(".", 1)[0] if "." in name else name
    stem = unicodedata.normalize("NFKD", stem).encode("ascii", "ignore").decode()
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_", stem).strip("_-")[:80] or "artifact"
    if re.fullmatch(r"(?i:CON|PRN|AUX|NUL|COM[0-9]|LPT[0-9])", stem):
        stem = "artifact_" + stem
    return stem + INGEST_EXTENSIONS[kind]


def _owner_hash(owner):
    _text(owner, 1024)
    if not owner.strip():
        _fail("Proprietário autenticado ausente.")
    return hashlib.sha256(b"iald-task-owner\0" + owner.encode("utf-8")).hexdigest()


def _uuid(value):
    if type(value) is not str:
        _fail("Identificador de artefato inválido.")
    try:
        parsed = uuid.UUID(value)
    except (ValueError, AttributeError):
        _fail("Identificador de artefato inválido.")
    if str(parsed) != value or parsed.version != 4 or parsed.variant != uuid.RFC_4122:
        _fail("Identificador de artefato inválido: use um UUID canônico.")
    return value


def _check_path(path, directory=None):
    """Inspect every ancestor without resolving links, including Windows junctions."""
    for part in reversed([path, *path.parents]):
        info = part.lstat()
        if stat.S_ISLNK(info.st_mode) or getattr(info, "st_file_attributes", 0) & 0x400:
            _fail("Caminho recusado: links simbólicos ou pontos de redirecionamento.")
        if part != path and not stat.S_ISDIR(info.st_mode):
            _fail("Caminho de armazenamento inválido.")
    if directory is True and not stat.S_ISDIR(info.st_mode):
        _fail("A pasta de armazenamento é inválida.")
    if directory is False and (not stat.S_ISREG(info.st_mode) or info.st_nlink != 1):
        _fail("Arquivo recusado: tipo especial ou vínculo físico.")
    return info


def _mkdir(path):
    _check_path(path.parent, True)
    try:
        path.mkdir(mode=0o700)
    except FileExistsError:
        pass
    _check_path(path, True)


def _read(path, limit):
    before = _check_path(path, False)
    if before.st_size > limit:
        _fail("Arquivo armazenado acima do limite permitido.")
    flags = os.O_RDONLY | getattr(os, "O_BINARY", 0) | getattr(os, "O_NOFOLLOW", 0) | getattr(os, "O_NONBLOCK", 0)
    fd = os.open(path, flags)
    with os.fdopen(fd, "rb") as handle:
        opened = os.fstat(handle.fileno())
        if (before.st_dev, before.st_ino) != (opened.st_dev, opened.st_ino) or opened.st_nlink != 1:
            _fail("O arquivo mudou durante a leitura.")
        raw = handle.read(limit + 1)
        after = _check_path(path, False)
        if (opened.st_dev, opened.st_ino, opened.st_size, opened.st_mtime_ns) != (
                after.st_dev, after.st_ino, after.st_size, after.st_mtime_ns):
            _fail("O arquivo mudou durante a leitura.")
    if len(raw) > limit or len(raw) != after.st_size:
        _fail("Tamanho do arquivo inválido.")
    return raw


def _write_new(path, raw):
    _check_path(path.parent, True)
    flags = os.O_WRONLY | os.O_CREAT | os.O_EXCL | getattr(os, "O_BINARY", 0) | getattr(os, "O_NOFOLLOW", 0)
    fd = os.open(path, flags, 0o600)
    with os.fdopen(fd, "wb") as handle:
        handle.write(raw)
        handle.flush()
        os.fsync(handle.fileno())
    if _check_path(path, False).st_size != len(raw):
        _fail("Falha na gravação completa do arquivo.")


def _sync_directory(path):
    if os.name != "nt":
        fd = os.open(path, os.O_RDONLY | getattr(os, "O_DIRECTORY", 0) | getattr(os, "O_NOFOLLOW", 0))
        try:
            os.fsync(fd)
        finally:
            os.close(fd)


def _cell_ref(value, rows, cols):
    match = re.fullmatch(r"\$?([A-Z]{1,3})\$?([1-9][0-9]{0,6})", value.upper())
    if not match:
        _fail("Referência de célula inválida.")
    column = 0
    for letter in match[1]:
        column = column * 26 + ord(letter) - 64
    row = int(match[2])
    if not (1 <= row <= rows and 1 <= column <= cols):
        _fail("A fórmula referencia uma célula fora da tabela permitida.")
    return row, column


_FORMULA_TOKEN = re.compile(r"\s*(\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6}(?![A-Za-z0-9_])|"
                            r"(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?|"
                            r"[A-Za-z_][A-Za-z0-9_.]*|[+*/^(),:\-])")


def _formula(value, rows, cols):
    """Parse a small non-evaluating arithmetic grammar; no Excel escape hatches."""
    _text(value, 1024)
    if not value.startswith("="):
        _fail("Fórmulas explícitas devem começar com =.")
    source = value[1:].strip()
    tokens, offset = [], 0
    while offset < len(source):
        match = _FORMULA_TOKEN.match(source, offset)
        if not match:
            _fail("Fórmula recusada: apenas aritmética e SUM/AVERAGE/MIN/MAX são permitidos.")
        tokens.append(match[1].upper())
        offset = match.end()
    if not tokens or len(tokens) > 256:
        _fail("Fórmula vazia ou muito complexa.")
    index = 0

    def peek():
        return tokens[index] if index < len(tokens) else ""

    def take(expected=None):
        nonlocal index
        token = peek()
        if not token or (expected is not None and token != expected):
            _fail("Sintaxe de fórmula inválida.")
        index += 1
        return token

    def expression(depth=0, minimum=0):
        if depth > 32:
            _fail("Fórmula com aninhamento excessivo.")
        token = take()
        is_range = False
        if token in ("+", "-"):
            is_range = expression(depth + 1, 4)
            if is_range:
                _fail("Intervalos só são permitidos como argumentos de agregação.")
        elif token == "(":
            is_range = expression(depth + 1)
            take(")")
        elif token in ("SUM", "AVERAGE", "MIN", "MAX"):
            take("(")
            expression(depth + 1)
            while peek() == ",":
                take(",")
                expression(depth + 1)
            take(")")
        elif re.fullmatch(r"\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6}", token):
            start = _cell_ref(token, rows, cols)
            if peek() == ":":
                take(":")
                end = _cell_ref(take(), rows, cols)
                if end[0] < start[0] or end[1] < start[1]:
                    _fail("Intervalo de fórmula invertido.")
                is_range = True
        elif re.fullmatch(r"(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:E[+-]?[0-9]+)?", token):
            if not math.isfinite(float(token)):
                _fail("Número não finito na fórmula.")
        else:
            _fail("Função ou referência não permitida na fórmula.")
        precedence = {"+": 1, "-": 1, "*": 2, "/": 2, "^": 3}
        while peek() in precedence and precedence[peek()] >= minimum:
            op = take()
            right_range = expression(depth + 1, precedence[op] + (op != "^"))
            if is_range or right_range:
                _fail("Intervalos só são permitidos como argumentos de agregação.")
        return is_range

    if expression() or index != len(tokens):
        _fail("Sintaxe de fórmula inválida.")
    return "=" + "".join(tokens)


def _scalar(value):
    if value is not None and type(value) not in (str, int, float, bool):
        _fail("Células devem conter somente texto, número, booleano ou nulo.")
    if type(value) is str:
        _text(value, 32767)
    if type(value) in (int, float) and (not math.isfinite(value) or abs(value) > 1e100):
        _fail("Número de célula fora do limite permitido.")
    return value


def _table(headers, rows):
    if type(headers) is not list or not 1 <= len(headers) <= MAX_COLUMNS:
        _fail("A tabela deve conter entre 1 e 100 colunas.")
    for header in headers:
        _text(header, 32767)
    if type(rows) is not list or len(rows) > MAX_ROWS or (len(rows) + 1) * len(headers) > MAX_CELLS:
        _fail("A tabela excede o limite de linhas ou células.")
    for row in rows:
        if type(row) is not list or len(row) != len(headers):
            _fail("Cada linha deve ter o mesmo número de células do cabeçalho.")
        for value in row:
            _scalar(value)


def _csv_bytes(headers, rows):
    _table(headers, rows)
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\r\n")
    for row in [headers, *rows]:
        writer.writerow(["'" + val if type(val) is str and val.lstrip(" \t\r\n\ufeff").startswith(("=", "+", "-", "@"))
                         else val for val in row])
    return output.getvalue().encode("utf-8")


def _spreadsheet(payload):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    _object(payload, ("sheets",))
    sheets = payload["sheets"]
    if type(sheets) is not list or not 1 <= len(sheets) <= MAX_SHEETS:
        _fail("A planilha deve conter entre 1 e 16 abas.")
    workbook = Workbook()
    workbook.remove(workbook.active)
    seen, total = set(), 0
    for spec in sheets:
        _object(spec, ("name", "headers", "rows"), ("formulas",))
        name = _text(spec["name"], 31)
        if not name.strip() or re.search(r"[\[\]:*?/\\]", name) or name.startswith("'") or name.endswith("'") or name.casefold() in seen:
            _fail("Nome de aba inválido ou duplicado.")
        seen.add(name.casefold())
        headers, rows = spec["headers"], spec["rows"]
        _table(headers, rows)
        total += (len(rows) + 1) * len(headers)
        if total > MAX_CELLS:
            _fail("A planilha excede 100000 células.")
        sheet = workbook.create_sheet(name)
        for r, values in enumerate([headers, *rows], 1):
            for c, value in enumerate(values, 1):
                cell = sheet.cell(r, c, value)
                if type(value) is str:
                    cell.data_type = "s"  # Never infer a formula from ordinary text.
                cell.font = Font(name="Calibri", size=11, color="172B3A")
        formulas = spec.get("formulas", {})
        if type(formulas) is not dict or len(formulas) > MAX_CELLS:
            _fail("O mapa de fórmulas é inválido.")
        targets = set()
        for target, formula in formulas.items():
            _text(target, 16)
            coordinates = _cell_ref(target, len(rows) + 1, len(headers))
            if coordinates[0] == 1 or coordinates in targets:
                _fail("Fórmula duplicada ou destinada ao cabeçalho.")
            if sheet.cell(*coordinates).value not in (None, ""):
                _fail("A fórmula substituiria um dado fornecido. Preserve os valores e reserve uma célula vazia para o resultado; cabeçalho na linha 1, dados a partir da linha 2.")
            targets.add(coordinates)
            sheet.cell(*coordinates, _formula(formula, len(rows) + 1, len(headers)))
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill("solid", fgColor="172B3A")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.column_dimensions[cell.column_letter].width = min(45, max(16, len(str(cell.value)) + 3))
        sheet.row_dimensions[1].height = 30
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _strip_template_parts(raw):
    """Remove the library template's unused thumbnail and printer binary."""
    from lxml import etree
    source, output = zipfile.ZipFile(io.BytesIO(raw)), io.BytesIO()
    with source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        removed = {n for n in source.namelist() if n.startswith("ppt/printerSettings/") or n.startswith("docProps/thumbnail.")}
        for name in source.namelist():
            if name in removed:
                continue
            content = source.read(name)
            if name.endswith(".rels") or name == "[Content_Types].xml":
                tree = etree.fromstring(content, etree.XMLParser(resolve_entities=False, no_network=True))
                for element in list(tree):
                    if element.get("PartName", "").lstrip("/") in removed or element.get("Extension") in ("bin", "jpeg"):
                        tree.remove(element)
                    elif element.tag.endswith("Relationship") and _relationship_path(name, element.get("Target", "")) in removed:
                        tree.remove(element)
                content = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(name, content)
    return output.getvalue()


def _slides(payload):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    _object(payload, ("slides",))
    specs = payload["slides"]
    if type(specs) is not list or not 1 <= len(specs) <= MAX_SLIDES:
        _fail("A apresentação deve conter entre 1 e 50 slides.")
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for spec in specs:
        _object(spec, ("title", "bullets"))
        title = _text(spec["title"], 90)
        bullets = spec["bullets"]
        if not title.strip() or type(bullets) is not list or len(bullets) > 6:
            _fail("Cada slide precisa de título e no máximo 6 tópicos.")
        for bullet in bullets:
            _text(bullet, 160)
            if "\n" in bullet or "\r" in bullet:
                _fail("Cada tópico deve ocupar um único parágrafo.")
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("142C35")
        heading = slide.shapes.title
        heading.left, heading.top, heading.width, heading.height = Inches(.75), Inches(.65), Inches(11.8), Inches(1.3)
        heading.text = title
        for paragraph in heading.text_frame.paragraphs:
            paragraph.font.name, paragraph.font.size, paragraph.font.bold = "Calibri", Pt(34), True
            paragraph.font.color.rgb = RGBColor.from_string("75DECC")
        body = slide.placeholders[1]
        body.left, body.top, body.width, body.height = Inches(.95), Inches(2.2), Inches(11.4), Inches(4.65)
        body.text_frame.clear()
        body.text_frame.word_wrap = True
        for index, bullet in enumerate(bullets):
            paragraph = body.text_frame.paragraphs[0] if index == 0 else body.text_frame.add_paragraph()
            paragraph.text, paragraph.level = bullet, 0
            paragraph.font.name, paragraph.font.size = "Calibri", Pt(22)
            paragraph.font.color.rgb = RGBColor.from_string("F3F8FA")
            paragraph.space_after = Pt(16)
    buffer = io.BytesIO()
    deck.save(buffer)
    return _strip_template_parts(buffer.getvalue())


_WEBSITE_CSS = "body{font-family:Arial,sans-serif;margin:0;background:#f5f7fa;color:#172b3a;line-height:1.6}main{max-width:960px;margin:48px auto;padding:32px;background:#ffffff}h1{font-size:36px}h2{color:#176b65}section{margin:32px 0}"
_CSS_PROPERTIES = frozenset("font-family font-size font-weight font-style line-height color background background-color margin margin-top margin-bottom margin-left margin-right padding padding-top padding-bottom padding-left padding-right max-width min-width width height border border-color border-width border-style border-radius text-align text-decoration display gap grid-template-columns letter-spacing white-space overflow-wrap".split())


def _safe_css(css):
    if len(css) > 20_000 or re.search(r"[^A-Za-z0-9\s#.,:%;{}>+*_\-]", css):
        _fail("CSS recusado: somente estilos locais simples são permitidos.")
    position = 0
    for match in re.finditer(r"\s*([A-Za-z0-9 .,#>:+*_\-]+)\{([^{}]*)\}", css):
        if match.start() != position:
            _fail("Estrutura CSS inválida.")
        for declaration in match[2].split(";"):
            if not declaration.strip():
                continue
            key, separator, val = declaration.partition(":")
            if key.strip().lower() not in _CSS_PROPERTIES or not separator or not re.fullmatch(r"[A-Za-z0-9\s#.,%_\-]+", val):
                _fail("Propriedade CSS não permitida.")
        position = match.end()
    if css[position:].strip():
        _fail("Estrutura CSS inválida.")
    return css


class _StaticHTML(HTMLParser):
    tags = frozenset("html head meta title style body main header footer section article div span h1 h2 h3 h4 h5 h6 p ul ol li strong b em i pre code blockquote br hr table thead tbody tr th td".split())
    voids = frozenset(("meta", "br", "hr"))

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.output, self.stack = [], []
        self.nodes = 0

    def handle_starttag(self, tag, attrs):
        self.nodes += 1
        if tag not in self.tags or self.nodes > 100_000 or len(self.stack) >= 64:
            _fail("HTML recusado: elemento ativo, externo ou complexo demais.")
        for key, value in attrs:
            if tag == "meta" and key == "charset" and value and value.lower() in ("utf-8", "utf8"):
                continue
            if key in ("class", "id", "lang") and value and re.fullmatch(r"[A-Za-z0-9 _-]{1,100}", value):
                continue
            _fail("HTML recusado: atributo não permitido.")
        self.output.append("<" + tag + "".join(' ' + k + '="' + html.escape(v, quote=True) + '"' for k, v in attrs) + ">")
        if tag not in self.voids:
            self.stack.append(tag)

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag not in self.voids:
            self.handle_endtag(tag)

    def handle_endtag(self, tag):
        if tag in self.voids or not self.stack or self.stack.pop() != tag:
            _fail("HTML inválido: fechamento de elemento inconsistente.")
        self.output.append("</" + tag + ">")

    def handle_data(self, data):
        self.output.append(_safe_css(data) if self.stack and self.stack[-1] == "style" else html.escape(data))

    def handle_decl(self, decl):
        if decl.lower() != "doctype html":
            _fail("Declaração HTML não permitida.")
        self.output.append("<!doctype html>")

    def handle_pi(self, data):
        _fail("Instrução de processamento HTML não permitida.")

    def unknown_decl(self, data):
        _fail("Declaração HTML não permitida.")


def _static_html(source):
    _text(source)
    parser = _StaticHTML()
    parser.feed(source)
    parser.close()
    if parser.stack or not parser.nodes:
        _fail("HTML inválido ou incompleto.")
    return "".join(parser.output).encode("utf-8")


def _website(payload):
    if "html" in payload:
        _object(payload, ("html",))
        return _static_html(payload["html"])
    _object(payload, ("title", "sections"))
    title = html.escape(_text(payload["title"], 200))
    sections = payload["sections"]
    if type(sections) is not list or not 1 <= len(sections) <= 100:
        _fail("O site deve conter entre 1 e 100 seções.")
    output = [f'<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><title>{title}</title><style>{_WEBSITE_CSS}</style></head><body><main><h1>{title}</h1>']
    for section in sections:
        _object(section, ("heading", "text"), ("bullets",))
        output.append("<section><h2>" + html.escape(_text(section["heading"], 200)) + "</h2><p>" + html.escape(_text(section["text"], 20_000)) + "</p>")
        bullets = section.get("bullets", [])
        if type(bullets) is not list or len(bullets) > 100:
            _fail("Lista de tópicos inválida.")
        if bullets:
            output.append("<ul>" + "".join("<li>" + html.escape(_text(b, 2000)) + "</li>" for b in bullets) + "</ul>")
        output.append("</section>")
    return _static_html("".join(output) + "</main></body></html>")


def _relationship_path(rels_name, target):
    if not target or any(c in target for c in "\\:%?#") or "\x00" in target:
        _fail("Relação Office externa ou inválida.")
    base = "" if rels_name == "_rels/.rels" else posixpath.dirname(posixpath.dirname(rels_name))
    path = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join(base, target))
    if path == ".." or path.startswith("../") or path.startswith("/"):
        _fail("Relação Office fora do pacote.")
    return path


def _office_zip(raw, kind):
    """Inspect every ZIP member and XML part before a library opens untrusted data."""
    common = r"(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/(?:app|core)\.xml)"
    xlsx = r"xl/(?:workbook\.xml|styles\.xml|sharedStrings\.xml|theme/theme[0-9]+\.xml|worksheets/sheet[0-9]+\.xml|_rels/workbook\.xml\.rels|worksheets/_rels/sheet[0-9]+\.xml\.rels)"
    pptx = r"ppt/(?:presentation\.xml|presProps\.xml|viewProps\.xml|tableStyles\.xml|theme/theme[0-9]+\.xml|_rels/presentation\.xml\.rels|(?:slides/(?:slide[0-9]+\.xml|_rels/slide[0-9]+\.xml\.rels)|slideLayouts/(?:slideLayout[0-9]+\.xml|_rels/slideLayout[0-9]+\.xml\.rels)|slideMasters/(?:slideMaster[0-9]+\.xml|_rels/slideMaster[0-9]+\.xml\.rels)))"
    pattern = common + "|" + (xlsx if kind == "spreadsheet" else pptx)
    required_main = "xl/workbook.xml" if kind == "spreadsheet" else "ppt/presentation.xml"
    trees, total, names, cells = {}, 0, set(), 0
    with zipfile.ZipFile(io.BytesIO(raw)) as package:
        members = package.infolist()
        if len(members) > MAX_ZIP_ENTRIES or not raw.startswith(b"PK\x03\x04"):
            _fail("Pacote Office inválido ou com entradas demais.")
        for info in members:
            name = info.filename
            total += info.file_size
            mode = info.external_attr >> 16
            if (name in names or not re.fullmatch(pattern, name) or info.is_dir()
                    or stat.S_IFMT(mode) not in (0, stat.S_IFREG) or info.flag_bits & 1
                    or info.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED)
                    or info.file_size > MAX_ZIP_MEMBER_BYTES or total > MAX_ZIP_EXPANDED_BYTES
                    or info.file_size > max(1, info.compress_size) * MAX_ZIP_RATIO):
                _fail("Pacote Office recusado: entrada insegura ou limite de descompressão.")
            names.add(name)
            with package.open(info) as member:
                content = member.read(MAX_ZIP_MEMBER_BYTES + 1)
            if len(content) != info.file_size:
                _fail("Pacote Office corrompido.")
            source = content.decode("utf-8-sig")
            if re.search(r"<!\s*(?:DOCTYPE|ENTITY)", source, re.I):
                _fail("XML Office recusado: entidades ou DTD não são permitidos.")
            tree = ET.fromstring(source)
            trees[name] = tree
            for element in tree.iter():
                tag = element.tag.rsplit("}", 1)[-1]
                if tag.lower() in {"externallink", "externalreference", "hyperlink", "hyperlinks", "hlinkclick", "hlinkhover", "oleobj", "oleobject", "control", "controls", "attachedtemplate"}:
                    _fail("Conteúdo Office ativo, vinculado ou não suportado.")
                if tag == "definedName":
                    # Excel's static filter range is not an executable named formula.
                    match = re.fullmatch(r"'(?:[^'\[\]!:*?/\\]|''){1,62}'!(\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6}):(\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6})", element.text or "")
                    if (kind != "spreadsheet" or element.get("name") != "_xlnm._FilterDatabase"
                            or not match or set(element.attrib) - {"name", "localSheetId", "hidden"}
                            or not re.fullmatch(r"[0-9]{1,2}", element.get("localSheetId", ""))):
                        _fail("Nomes definidos executáveis ou externos não são permitidos.")
                    _cell_ref(match[1], MAX_ROWS + 1, MAX_COLUMNS)
                    _cell_ref(match[2], MAX_ROWS + 1, MAX_COLUMNS)
                if tag == "extLst":
                    # The stock python-pptx template contains these inert numeric hints.
                    namespace = "{http://schemas.microsoft.com/office/powerpoint/2010/main}"
                    hints = {namespace + hint for hint in ("discardImageEditData", "defaultImageDpi", "creationId")}
                    if kind != "slides" or len(element) > 3:
                        _fail("Extensão Office não permitida.")
                    for extension in element:
                        if (extension.tag != "{http://schemas.openxmlformats.org/presentationml/2006/main}ext"
                                or len(extension) != 1 or extension[0].tag not in hints
                                or set(extension[0].attrib) != {"val"} or len(extension[0])
                                or not re.fullmatch(r"[0-9]{1,10}", extension[0].get("val", ""))):
                            _fail("Extensão Office não permitida.")
                if tag == "f" and kind == "spreadsheet":
                    if element.attrib:
                        _fail("Fórmulas compartilhadas, matriciais ou especiais não são permitidas.")
                    _formula("=" + (element.text or ""), MAX_ROWS + 1, MAX_COLUMNS)
                if tag == "c" and name.startswith("xl/worksheets/"):
                    _cell_ref(element.get("r", ""), MAX_ROWS + 1, MAX_COLUMNS)
                    cells += 1
                    if cells > MAX_CELLS:
                        _fail("A planilha excede o limite de células.")
                if tag == "dimension" and name.startswith("xl/worksheets/"):
                    for endpoint in element.get("ref", "").split(":"):
                        _cell_ref(endpoint, MAX_ROWS + 1, MAX_COLUMNS)
        if not {"[Content_Types].xml", "_rels/.rels", required_main} <= names:
            _fail("Pacote Office incompleto ou de formato incorreto.")
        rel_prefix = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
        allowed_relationships = {rel_prefix + s for s in ("officeDocument", "extended-properties", "worksheet", "styles", "theme", "sharedStrings", "slide", "slideMaster", "slideLayout", "presProps", "viewProps", "tableStyles")}
        allowed_relationships.add("http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties")
        main_type = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml" if kind == "spreadsheet" else "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml")
        main_declared = False
        for name, tree in trees.items():
            if name.endswith(".rels"):
                ids = set()
                for rel in tree:
                    if (rel.tag != "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"
                            or rel.get("TargetMode", "Internal") != "Internal"
                            or rel.get("Type") not in allowed_relationships or not rel.get("Id") or rel.get("Id") in ids
                            or _relationship_path(name, rel.get("Target", "")) not in names):
                        _fail("Relação Office externa, duplicada ou não permitida.")
                    ids.add(rel.get("Id"))
            if name == "[Content_Types].xml":
                for entry in tree:
                    content_type = entry.get("ContentType", "")
                    if any(token in content_type.lower() for token in ("macro", "vba", "ole", "activex", "javascript")) or not content_type.endswith("xml"):
                        _fail("Tipo de conteúdo Office não permitido.")
                    if entry.get("PartName") == "/" + required_main and content_type == main_type:
                        main_declared = True
        if not main_declared:
            _fail("O tipo declarado do documento Office é inválido.")


def _validate_office(raw, kind):
    _office_zip(raw, kind)
    if kind == "spreadsheet":
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(raw), data_only=False, keep_links=False)
        try:
            if not 1 <= len(workbook.worksheets) <= MAX_SHEETS:
                _fail("Número de abas fora do limite permitido.")
            total = 0
            for sheet in workbook.worksheets:
                rows, cols = sheet.max_row, sheet.max_column
                total += rows * cols
                if rows > MAX_ROWS + 1 or cols > MAX_COLUMNS or total > MAX_CELLS:
                    _fail("Dimensões da planilha acima do limite permitido.")
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.hyperlink:
                            _fail("Hiperlinks não são permitidos na planilha.")
                        if cell.data_type == "f":
                            _formula(cell.value, rows, cols)
                        else:
                            _scalar(cell.value)
        finally:
            workbook.close()
    else:
        from pptx import Presentation
        deck = Presentation(io.BytesIO(raw))
        if not 1 <= len(deck.slides) <= MAX_SLIDES:
            _fail("Número de slides fora do limite permitido.")
        for slide in deck.slides:
            if len(slide.shapes) > 100:
                _fail("Slide complexo demais para este formato.")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    _text(shape.text, 20_000)


def _generate(kind, payload):
    if kind in ("python", "text"):
        _object(payload, ("content",))
        return _text(payload["content"]).encode("utf-8")
    if kind == "json":
        _object(payload, ("data",))
        return (json.dumps(payload["data"], ensure_ascii=False, allow_nan=False, indent=2) + "\n").encode("utf-8")
    if kind == "csv":
        _object(payload, ("headers", "rows"))
        return _csv_bytes(payload["headers"], payload["rows"])
    return {"spreadsheet": _spreadsheet, "slides": _slides, "website": _website}[kind](payload)


def _validate_png(raw):
    from PIL import Image
    if not raw.startswith(b"\x89PNG\r\n\x1a\n"):
        _fail("PNG inválido: assinatura ausente.")
    offset, seen, ended, chunks = 8, set(), False, 0
    allowed = {b"IHDR", b"PLTE", b"IDAT", b"IEND", b"tRNS", b"sRGB", b"gAMA", b"cHRM", b"pHYs"}
    while offset < len(raw):
        chunks += 1
        if chunks > 2048 or offset + 12 > len(raw):
            _fail("PNG inválido ou complexo demais.")
        length = struct.unpack_from(">I", raw, offset)[0]
        tag = raw[offset + 4:offset + 8]
        end = offset + 12 + length
        if tag not in allowed or end > len(raw) or ended:
            _fail("PNG recusado: bloco inválido, ativo ou não suportado.")
        data = raw[offset + 8:end - 4]
        if zlib.crc32(tag + data) & 0xFFFFFFFF != struct.unpack_from(">I", raw, end - 4)[0]:
            _fail("PNG corrompido: CRC divergente.")
        if not seen and tag != b"IHDR":
            _fail("PNG sem cabeçalho inicial.")
        if tag == b"IHDR":
            if seen or length != 13:
                _fail("Cabeçalho PNG inválido.")
            width, height = struct.unpack_from(">II", data)
            if not width or not height or width > 4096 or height > 4096 or width * height > 4_000_000:
                _fail("PNG excede o limite de 4 milhões de pixels ou 4096 por dimensão.")
        if tag == b"IEND":
            if length or b"IDAT" not in seen:
                _fail("Final PNG inválido.")
            ended = True
        seen.add(tag)
        offset = end
    if not ended:
        _fail("PNG incompleto.")
    with Image.open(io.BytesIO(raw)) as picture:
        if picture.format != "PNG" or getattr(picture, "n_frames", 1) != 1:
            _fail("Somente PNG estático é permitido.")
        picture.verify()
    with Image.open(io.BytesIO(raw)) as picture:
        picture.load()


def _validate_pdf(raw):
    """Bounded classic PDFs only; never decode content streams or run actions."""
    from pypdf import PdfReader
    from pypdf.generic import IndirectObject, DictionaryObject, ArrayObject
    from pypdf.errors import PdfReadError
    if not re.match(br"%PDF-1\.[0-7](?:\r|\n)", raw) or not raw.rstrip().endswith(b"%%EOF"):
        _fail("PDF inválido ou versão não suportada.")
    # Reject compressed object/xref streams BEFORE the parser can inflate them.
    normalized = re.sub(br"#([0-9a-fA-F]{2})", lambda m: bytes([int(m[1], 16)]), raw)
    if re.search(br"/(?:ObjStm|XRef|Encrypt)\b", normalized) or normalized.count(b"%%EOF") != 1:
        _fail("PDF recusado: criptografia, objetos comprimidos ou revisões incrementais.")
    blocked = {"/JS", "/JavaScript", "/OpenAction", "/AA", "/A", "/Launch", "/URI", "/GoToR", "/GoToE",
               "/EmbeddedFiles", "/EmbeddedFile", "/Filespec", "/EF", "/AF", "/AcroForm", "/XFA",
               "/RichMedia", "/Movie", "/Sound", "/SubmitForm", "/ImportData", "/Collection",
               "/F", "/FFilter", "/FDecodeParms", "/Ref", "/3D", "/Rendition"}
    try:
        reader = PdfReader(io.BytesIO(raw), strict=True)
        if reader.is_encrypted or sum(len(entries) for entries in reader.xref.values()) > 5000:
            _fail("PDF criptografado ou com objetos demais.")
        visited, count = set(), [0]

        def visit(value, depth=0):
            count[0] += 1
            if count[0] > 100_000 or depth > 64:
                _fail("PDF complexo demais.")
            if isinstance(value, IndirectObject):
                identity = (value.idnum, value.generation)
                if identity in visited:
                    return
                visited.add(identity)
                return visit(value.get_object(), depth + 1)
            if isinstance(value, DictionaryObject):
                for key, child in value.items():
                    if str(key) in blocked or (isinstance(child, str) and child in blocked):
                        _fail("PDF recusado: ações, formulários, anexos ou recursos externos.")
                    visit(child, depth + 1)
            elif isinstance(value, ArrayObject):
                for child in value:
                    visit(child, depth + 1)

        visit(reader.trailer)
        for generation, entries in reader.xref.items():
            for object_id in entries:
                if object_id and generation != 65535:
                    visit(IndirectObject(object_id, generation, reader))
        if not 1 <= len(reader.pages) <= 200:
            _fail("O PDF deve conter entre 1 e 200 páginas.")
    except PdfReadError as exc:
        raise ArtifactError("PDF inválido ou corrompido.") from exc


def _utc_timestamp(epoch):
    return datetime.fromtimestamp(epoch, timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


class ArtifactStore:
    @_public
    def __init__(self, root):
        root = Path(root)
        if ".." in root.parts:
            _fail("A raiz de armazenamento não pode conter navegação para pasta superior.")
        self.root = Path(os.path.abspath(root))
        _check_path(self.root, True)
        runtime = self.root / ".iald_runtime"
        _mkdir(runtime)
        # The app owns this optional, explicitly provisioned dependency directory.
        # Appending provides a fallback without upgrading or shadowing installed libs.
        vendor = runtime / "task_dependencies"
        if os.path.lexists(vendor):
            _check_path(vendor, True)
            if str(vendor) not in sys.path:
                sys.path.append(str(vendor))
        self.storage = runtime / "task_artifacts"
        _mkdir(self.storage)
        self.tickets = self.storage / ".tickets"
        _mkdir(self.tickets)

    def _owner_directory(self, owner, create=False):
        digest = _owner_hash(owner)
        _check_path(self.storage, True)
        path = self.storage / digest
        if create:
            _mkdir(path)
        else:
            try:
                _check_path(path, True)
            except FileNotFoundError:
                _fail("Artefato não encontrado para este proprietário.")
        return path, digest

    @contextmanager
    def _lock(self, directory):
        path = directory / ".write.lock"
        try:
            _write_new(path, b"artifact-write\n")
        except FileExistsError:
            _fail("Armazenamento ocupado ou gravação anterior pendente de recuperação.")
        try:
            yield
        finally:
            _check_path(path, False)
            path.unlink()

    def _record(self, directory, owner_digest, artifact_id):
        _uuid(artifact_id)
        path = directory / artifact_id
        try:
            _check_path(path, True)
            metadata = _load_json(_read(path / "metadata.json", MAX_METADATA_BYTES))
        except FileNotFoundError:
            _fail("Artefato não encontrado para este proprietário.")
        expected = {"schema_version", "id", "owner_hash", "kind", "filename", "sha256", "bytes", "url", "status"}
        if type(metadata) is not dict or set(metadata) != expected:
            _fail("Metadados do artefato inválidos.")
        kind = metadata["kind"]
        if type(kind) is not str or kind not in INGEST_EXTENSIONS:
            _fail("Formato armazenado inválido.")
        filename = metadata["filename"]
        if (type(metadata["schema_version"]) is not int or metadata["schema_version"] != 1
                or metadata["id"] != artifact_id or metadata["owner_hash"] != owner_digest
                or filename != _filename(filename, kind) or metadata["url"] != _artifact_url(artifact_id)
                or metadata["status"] != "generated" or type(metadata["bytes"]) is not int
                or not 0 <= metadata["bytes"] <= MAX_ARTIFACT_BYTES
                or type(metadata["sha256"]) is not str or not re.fullmatch(r"[0-9a-f]{64}", metadata["sha256"])):
            _fail("Metadados adulterados ou inconsistentes.")
        if {item.name for item in path.iterdir()} != {"metadata.json", filename}:
            _fail("Conteúdo armazenado inesperado ou incompleto.")
        content_path = path / filename
        raw = _read(content_path, MAX_ARTIFACT_BYTES)
        if len(raw) != metadata["bytes"] or hashlib.sha256(raw).hexdigest() != metadata["sha256"]:
            _fail("Integridade do artefato inválida: tamanho ou SHA-256 divergente.")
        return content_path, metadata

    def _usage(self, directory, digest):
        count, total = 0, 0
        for entry in directory.iterdir():
            if entry.name == ".write.lock":
                continue
            # Unexpected files and crash leftovers block further writes.
            _, metadata = self._record(directory, digest, entry.name)
            count += 1
            total += metadata["bytes"]
        return count, total

    def _store(self, owner, kind, filename, raw):
        if len(raw) > MAX_ARTIFACT_BYTES:
            _fail("O artefato excede o limite de 20 MiB.")
        if kind in ("spreadsheet", "slides"):
            _validate_office(raw, kind)
        directory, digest = self._owner_directory(owner, create=True)
        with self._lock(directory):
            count, total = self._usage(directory, digest)
            if count >= MAX_OWNER_ARTIFACTS or total + len(raw) > MAX_OWNER_BYTES:
                _fail("Cota do proprietário excedida: máximo de 5000 arquivos e 2 GiB.")
            artifact_id = str(uuid.uuid4())
            final = directory / artifact_id
            pending = directory / (".pending_" + artifact_id)
            metadata = {"schema_version": 1, "id": artifact_id, "owner_hash": digest, "kind": kind,
                        "filename": filename, "sha256": hashlib.sha256(raw).hexdigest(), "bytes": len(raw),
                        "url": _artifact_url(artifact_id), "status": "generated"}
            encoded = json.dumps(metadata, ensure_ascii=True, sort_keys=True).encode("utf-8")
            if len(encoded) > MAX_METADATA_BYTES:
                _fail("Metadados acima do limite permitido.")
            _check_path(directory, True)
            pending.mkdir(mode=0o700)  # Exclusive; never reuse an uncertain identity.
            try:
                _write_new(pending / filename, raw)
                _write_new(pending / "metadata.json", encoded)
                if _read(pending / filename, MAX_ARTIFACT_BYTES) != raw or _read(pending / "metadata.json", MAX_METADATA_BYTES) != encoded:
                    _fail("Falha ao conferir a gravação do artefato.")
                _sync_directory(pending)
                _check_path(directory, True)
                if os.path.lexists(final):
                    _fail("Identificador já utilizado; gravação recusada.")
                os.rename(pending, final)  # Publish bytes and metadata together.
                _sync_directory(directory)
            except BaseException:
                if pending.exists():
                    _check_path(pending, True)
                    for child in pending.iterdir():
                        _check_path(child, False)
                        child.unlink()
                    pending.rmdir()
                raise
            ticket = self._issue_ticket(directory, digest, artifact_id)
        public = {key: metadata[key] for key in ("id", "filename", "sha256", "bytes", "url")}
        public.update(ticket)
        return public

    @_public
    def create(self, owner, kind, name, payload):
        _owner_hash(owner)
        if type(kind) is not str or kind not in EXTENSIONS:
            _fail("Formato de artefato não permitido.")
        filename = _filename(name, kind)
        if type(payload) is not dict:
            _fail("O conteúdo do artefato deve ser um objeto JSON.")
        _json_tree(payload)
        artifact = self._store(owner, kind, filename, _generate(kind, payload))
        return {"status": "generated", "artifacts": [artifact]}

    @_public
    def ingest(self, owner, filename, raw):
        """Return a single artifact dict. CSV/HTML normalize; other bytes stay intact."""
        _owner_hash(owner)
        _text(filename, 255)
        suffix = Path(filename).suffix.lower()
        kind = {extension: name for name, extension in INGEST_EXTENSIONS.items()}.get(suffix)
        if suffix == ".htm":
            kind = "website"
        if kind is None:
            _fail("Extensão não permitida para arquivos do provedor.")
        filename = _filename(filename, kind)
        if type(raw) is not bytes or len(raw) > MAX_ARTIFACT_BYTES:
            _fail("Forneça bytes de um arquivo de até 20 MiB.")
        if kind == "png":
            _validate_png(raw)
        elif kind == "pdf":
            _validate_pdf(raw)
        elif kind not in ("spreadsheet", "slides"):
            source = _text(raw.decode("utf-8-sig"))
            if kind == "json":
                _load_json(source)
            elif kind == "csv":
                rows = []
                for row in csv.reader(io.StringIO(source, newline=""), strict=True):
                    if len(rows) >= MAX_ROWS + 1 or len(row) > MAX_COLUMNS:
                        _fail("CSV excede os limites permitidos.")
                    rows.append(row)
                if not rows:
                    _fail("CSV vazio: informe um cabeçalho.")
                raw = _csv_bytes(rows[0], rows[1:])
            elif kind == "website":
                raw = _static_html(source)
        return self._store(owner, kind, filename, raw)

    @_public
    def resolve(self, owner, artifact_id):
        _uuid(artifact_id)
        directory, digest = self._owner_directory(owner)
        path, metadata = self._record(directory, digest, artifact_id)
        return path, self._download_metadata(metadata)

    @staticmethod
    def _download_metadata(metadata):
        metadata = dict(metadata)
        metadata["mime"] = MIME_TYPES[metadata["kind"]]
        metadata["download_headers"] = {
            "Content-Type": "application/octet-stream",
            "Content-Disposition": 'attachment; filename="' + metadata["filename"] + '"',
            "Content-Length": str(metadata["bytes"]),
            "X-Content-Type-Options": "nosniff",
            "Content-Security-Policy": "sandbox; default-src 'none'; frame-ancestors 'none'",
            "Cache-Control": "private, no-store",
            "Referrer-Policy": "no-referrer",
        }
        return metadata

    def _issue_ticket(self, directory, digest, artifact_id):
        # Called only while holding this owner's write lock.
        self._record(directory, digest, artifact_id)
        _check_path(self.tickets, True)
        token = secrets.token_urlsafe(32)
        issued = int(time.time())
        expires = issued + TICKET_TTL_SECONDS
        record = {"schema_version": 1, "id": artifact_id, "owner_hash": digest,
                  "token_sha256": hashlib.sha256(token.encode("ascii")).hexdigest(),
                  "issued_at": issued, "expires_at": expires}
        raw = json.dumps(record, sort_keys=True).encode("ascii")
        temporary = self.tickets / (".pending_" + str(uuid.uuid4()))
        destination = self.tickets / (artifact_id + ".json")
        try:
            _write_new(temporary, raw)
            if os.path.lexists(destination):
                _check_path(destination, False)
            _check_path(self.tickets, True)
            os.replace(temporary, destination)
            _sync_directory(self.tickets)
        finally:
            if os.path.lexists(temporary):
                _check_path(temporary, False)
                temporary.unlink()
        return {"mobile_url": PUBLIC_BASE_URL + "/task-download/" + artifact_id + "/" + token,
                "expires_at": _utc_timestamp(expires)}

    @_public
    def issue_ticket(self, owner, artifact_id, ttl=TICKET_TTL_SECONDS):
        """Return a new URL; lifetime is fixed and older tickets are revoked."""
        if type(ttl) is not int or ttl != TICKET_TTL_SECONDS:
            _fail("A validade do link temporário é fixa: 900 segundos.")
        _uuid(artifact_id)
        directory, digest = self._owner_directory(owner)
        with self._lock(directory):
            return self._issue_ticket(directory, digest, artifact_id)["mobile_url"]

    @_public
    def resolve_ticket(self, artifact_id, token):
        """Resolve only after validating an opaque, artifact-specific capability."""
        _uuid(artifact_id)
        error = "Link temporário inválido ou expirado."
        if type(token) is not str or not re.fullmatch(r"[A-Za-z0-9_-]{43}", token):
            _fail(error)
        try:
            record = _load_json(_read(self.tickets / (artifact_id + ".json"), MAX_METADATA_BYTES))
        except (ArtifactError, OSError, ValueError):
            _fail(error)
        fields = {"schema_version", "id", "owner_hash", "token_sha256", "issued_at", "expires_at"}
        now = int(time.time())
        if (type(record) is not dict or set(record) != fields or type(record["schema_version"]) is not int
                or record["schema_version"] != 1 or record["id"] != artifact_id
                or type(record["owner_hash"]) is not str or not re.fullmatch(r"[0-9a-f]{64}", record["owner_hash"])
                or type(record["token_sha256"]) is not str or not re.fullmatch(r"[0-9a-f]{64}", record["token_sha256"])
                or type(record["issued_at"]) is not int or type(record["expires_at"]) is not int
                or record["expires_at"] - record["issued_at"] != TICKET_TTL_SECONDS
                or not record["issued_at"] <= now < record["expires_at"]):
            _fail(error)
        candidate = hashlib.sha256(token.encode("ascii")).hexdigest()
        if not hmac.compare_digest(candidate, record["token_sha256"]):
            _fail(error)
        # Do not open ANY owner's artifact before the capability check succeeds.
        directory = self.storage / record["owner_hash"]
        path, metadata = self._record(directory, record["owner_hash"], artifact_id)
        metadata = self._download_metadata(metadata)
        metadata["expires_at"] = _utc_timestamp(record["expires_at"])
        return path, metadata
