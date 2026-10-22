"""Read only metadata/results of this new task pilot, never credentials."""
from pathlib import Path
import json,sys,hashlib,sqlite3
ROOT=Path(r'C:\IALD\projetos_pyhton\IALD')
sys.path.insert(0,str(ROOT))
from iald_task_artifacts import ArtifactStore
store=ArtifactStore(ROOT)
report={'tasks':[],'artifacts':[],'providers':[]}
for p in sorted((ROOT/'.iald_runtime/task_runs').glob('*.json'),key=lambda p:p.stat().st_mtime)[-5:]:
    row=json.loads(p.read_bytes())
    report['tasks'].append({'id':row['id'],'created':row['created'],'steps':list(row['steps'].values())})
folder,_=store._owner_directory('luiz')
for p in folder.glob('*/metadata.json'):
    meta=json.loads(p.read_bytes());file,checked=store.resolve('luiz',meta['id'])
    item={k:meta[k] for k in ('id','filename','bytes','sha256')}
    if file.suffix=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(file,read_only=True,data_only=False)
        item['sheets']={s.title:list(s.values) for s in book};book.close()
    elif file.suffix=='.pptx':
        from pptx import Presentation
        item['slides']=[[s.text for s in slide.shapes if s.has_text_frame] for slide in Presentation(file).slides]
    elif meta['filename']=='execution.json':
        r=json.loads(file.read_bytes())
        item['execution']={k:r.get(k) for k in ('execution_id','revision','state','evidence_status','success','error','http_status_code','dispatch_attempted','usage','input_tokens','reserved_total_tokens')}
        item['execution']['results']=r.get('results',[])
    report['artifacts'].append(item)
db=sqlite3.connect((ROOT/'.iald_runtime/app_evaluation/metrics.sqlite3').as_uri()+'?mode=ro',uri=True)
for (raw,) in db.execute('SELECT data FROM turns ORDER BY rowid DESC LIMIT 5'):
    row=json.loads(raw)
    report['providers'].append({k:row.get(k) for k in ('id','provider','status','error_code','utc','elapsed_seconds','calls')})
db.close()
Path(__file__).with_name('runtime_observation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
