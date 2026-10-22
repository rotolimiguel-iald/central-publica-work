"""Offline helper tests. Run with an installed-dependency Python and -B.

Only temporary local fixtures are created. No production imports, network,
subprocesses, generated scripts, formula engines or renderers are invoked.
The application's isolated vendor path can be placed on PYTHONPATH by the caller.
"""

import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import socket
import stat
import sys
import tempfile
import unittest
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch

sys.dont_write_bytecode = True
MODULE_PATH = Path(__file__).parent / "stage" / "iald_task_artifacts.py"
SPEC = importlib.util.spec_from_file_location("iald_task_artifacts", MODULE_PATH)
art = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(art)


def rewrite_zip(raw, updates=None, extra=None):
    updates, extra = updates or {}, extra or {}
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename in updates:
                data = updates[entry.filename](data)
            target.writestr(entry, data)
        for name, data in extra.items():
            target.writestr(name, data)
    return output.getvalue()


class ArtifactTests(unittest.TestCase):
    def test_formula_cannot_replace_a_supplied_value(self):
        payload={'sheets':[{'name':'Numbers','headers':['Item','Value'],
            'rows':[['A',2],['B',4],['C',6],['Total',None]],'formulas':{'B4':'=SUM(B2:B3)'}}]}
        with self.assertRaisesRegex(art.ArtifactError,'substituiria'):
            self.store.create(self.owner,'spreadsheet','numbers',payload)
        payload['sheets'][0]['formulas']={'B5':'=SUM(B2:B4)'}
        ref=self.store.create(self.owner,'spreadsheet','numbers',payload)['artifacts'][0]
        file,_=self.store.resolve(self.owner,ref['id'])
        from openpyxl import load_workbook
        book=load_workbook(file);self.assertEqual(book.active['B4'].value,6)
        self.assertEqual(book.active['B5'].value,'=SUM(B2:B4)');book.close()

    @classmethod
    def setUpClass(cls):
        cls.network_guard = patch.object(socket, "socket", side_effect=AssertionError("Network is forbidden"))
        cls.network_guard.start()
        cls.xlsx_bytes = art._spreadsheet({"sheets": [{"name": "Data", "headers": ["Label", "Value"],
            "rows": [["a", 2], ["b", 3], ["total", None]], "formulas": {"B4": "=SUM(B2:B3)"}}]})
        cls.pptx_bytes = art._slides({"slides": [{"title": "Local title", "bullets": ["First item", "Second item"]}]})

    @classmethod
    def tearDownClass(cls):
        cls.network_guard.stop()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="iald-artifact-test-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.store = art.ArtifactStore(self.root)
        self.owner = "session-owner@example.test"

    def create(self, kind="text", name="report", payload=None, owner=None):
        result = self.store.create(owner or self.owner, kind, name, payload or {"content": "hello"})
        self.assertEqual(result["status"], "generated")
        self.assertEqual(len(result["artifacts"]), 1)
        return result["artifacts"][0]

    def resolve(self, artifact):
        return self.store.resolve(self.owner, artifact["id"])

    def test_contract_and_fixed_urls(self):
        record = self.create()
        self.assertEqual(set(record), {"id", "filename", "sha256", "bytes", "url", "mobile_url", "expires_at"})
        self.assertEqual(record["url"], "https://demo.iald.ia.br/api/task-artifacts/" + record["id"])
        self.assertTrue(record["mobile_url"].startswith("https://demo.iald.ia.br/task-download/"))
        self.assertTrue(record["expires_at"].endswith("Z"))
        self.assertNotIn(str(self.root), json.dumps(record))
        path, meta = self.resolve(record)
        self.assertEqual(path.read_bytes(), b"hello")
        self.assertEqual(meta["mime"], "text/plain")
        self.assertEqual(meta["bytes"], 5)
        self.assertEqual(meta["sha256"], hashlib.sha256(b"hello").hexdigest())
        self.assertNotIn(self.owner, path.as_posix())
        self.assertEqual(len(path.parent.parent.name), 64)

    def test_ingest_returns_single_dictionary_and_markdown_mime(self):
        record = self.store.ingest(self.owner, "notes.md", b"# Notes\n")
        self.assertNotIn("artifacts", record)
        self.assertEqual(self.resolve(record)[1]["mime"], "text/markdown")

    def test_tool_contract(self):
        tool = art.TOOL_DEFINITION
        self.assertEqual(tool["function"]["name"], "create_artifact")
        parameters = tool["function"]["parameters"]
        self.assertEqual(set(parameters["properties"]), {"kind", "name", "payload"})
        self.assertEqual(set(parameters["required"]), {"kind", "name", "payload"})
        self.assertEqual(set(parameters["properties"]["kind"]["enum"]), set(art.EXTENSIONS))
        self.assertIn("mobile_url", tool["function"]["description"])

    def test_explicit_root_dependency_fallback(self):
        vendor = self.root / ".iald_runtime" / "task_dependencies"
        vendor.mkdir()
        with patch.object(sys, "path", list(sys.path)):
            art.ArtifactStore(self.root)
            art.ArtifactStore(self.root)
            self.assertEqual(sys.path.count(str(vendor)), 1)

    def test_ownership_and_restart(self):
        record = self.create()
        other = art.ArtifactStore(self.root)
        self.assertEqual(other.resolve(self.owner, record["id"])[0].read_bytes(), b"hello")
        with self.assertRaises(art.ArtifactError):
            other.resolve("different-owner", record["id"])
        self.assertFalse((self.store.storage / art._owner_hash("different-owner")).exists())

    def test_invalid_owner_and_uuid(self):
        for owner in ("", "  ", None, {}, "bad\0owner"):
            with self.subTest(owner=owner), self.assertRaises(art.ArtifactError):
                self.store.create(owner, "text", "x", {"content": "x"})
        for identity in ("../file", str(uuid.uuid1()), "123", str(uuid.uuid4()).upper(), {}, "a" * 64):
            with self.subTest(identity=identity), self.assertRaises(art.ArtifactError):
                self.store.resolve(self.owner, identity)

    def test_path_traversal_and_names(self):
        for name in ("../a", "..\\a", "C:\\tmp\\a", "a/b", "a:stream", "..", "a\r\nHeader"):
            with self.subTest(name=name), self.assertRaises(art.ArtifactError):
                self.create(name=name)
        record = self.create(name="CON.txt")
        self.assertEqual(record["filename"], "artifact_CON.txt")
        self.assertEqual(self.create(name="Relatório final!.txt")["filename"], "Relatorio_final.txt")

    def test_tampered_bytes(self):
        record = self.create()
        path, _ = self.resolve(record)
        path.write_bytes(b"hallo")
        with self.assertRaisesRegex(art.ArtifactError, "SHA-256"):
            self.resolve(record)

    def test_tampered_metadata_fields(self):
        record = self.create()
        path, _ = self.resolve(record)
        metadata_path = path.parent / "metadata.json"
        original = metadata_path.read_bytes()
        for key, value in (("owner_hash", "0" * 64), ("filename", "../../escape.txt"), ("bytes", True),
                           ("url", "https://evil.test/"), ("id", str(uuid.uuid4())), ("status", "tests_passed"),
                           ("sha256", "0" * 64), ("schema_version", True), ("kind", "exe")):
            metadata = json.loads(original)
            metadata[key] = value
            metadata_path.write_text(json.dumps(metadata), encoding="utf-8")
            with self.subTest(key=key), self.assertRaises(art.ArtifactError):
                self.resolve(record)
            metadata_path.write_bytes(original)

    def test_duplicate_metadata_keys_and_extra_files(self):
        record = self.create()
        path, _ = self.resolve(record)
        meta_path = path.parent / "metadata.json"
        original = meta_path.read_bytes()
        meta_path.write_bytes(original[:-1] + b',"bytes":5}')
        with self.assertRaises(art.ArtifactError):
            self.resolve(record)
        meta_path.write_bytes(original)
        (path.parent / "extra.txt").write_bytes(b"extra")
        with self.assertRaises(art.ArtifactError):
            self.resolve(record)

    def test_hardlink_rejected(self):
        record = self.create()
        path, _ = self.resolve(record)
        alias = self.root / "alias"
        try:
            os.link(path, alias)
        except OSError as exc:
            self.skipTest("Host cannot create a hardlink: " + str(exc))
        with self.assertRaises(art.ArtifactError):
            self.resolve(record)

    def test_symlink_root_and_file_rejected(self):
        outside = self.root / "outside"
        outside.mkdir()
        link = self.root / "link"
        try:
            link.symlink_to(outside, target_is_directory=True)
        except OSError as exc:
            self.skipTest("Host does not grant symlink privilege: " + str(exc))
        with self.assertRaises(art.ArtifactError):
            art.ArtifactStore(link)
        record = self.create()
        path, _ = self.resolve(record)
        target = outside / "payload"
        target.write_bytes(b"hello")
        path.unlink()
        path.symlink_to(target)
        with self.assertRaises(art.ArtifactError):
            self.resolve(record)

    def test_reparse_point_rejected_with_mock(self):
        original = Path.lstat
        def fake(path, *args, **kwargs):
            info = original(path, *args, **kwargs)
            if path == self.store.storage:
                from types import SimpleNamespace
                return SimpleNamespace(st_mode=stat.S_IFDIR, st_file_attributes=0x400)
            return info
        with patch.object(Path, "lstat", fake), self.assertRaises(art.ArtifactError):
            self.create()

    def test_artifact_and_owner_quotas(self):
        with patch.object(art, "MAX_ARTIFACT_BYTES", 4), self.assertRaises(art.ArtifactError):
            self.create()
        self.create()
        with patch.object(art, "MAX_OWNER_ARTIFACTS", 1), self.assertRaises(art.ArtifactError):
            self.create()
        with patch.object(art, "MAX_OWNER_BYTES", 8), self.assertRaises(art.ArtifactError):
            self.create()
        self.create(owner="another-owner")

    def test_lock_blocks_and_atomic_failure_cleans_pending(self):
        directory, _ = self.store._owner_directory(self.owner, True)
        lock = directory / ".write.lock"
        lock.write_bytes(b"busy")
        with self.assertRaises(art.ArtifactError):
            self.create()
        self.assertEqual(lock.read_bytes(), b"busy")
        lock.unlink()
        original = art._write_new
        def failing(path, raw):
            if path.name == "metadata.json":
                raise OSError("simulated disk error")
            return original(path, raw)
        with patch.object(art, "_write_new", failing), self.assertRaises(art.ArtifactError):
            self.create()
        self.assertEqual(list(directory.iterdir()), [])
        self.create()

    def test_publish_is_one_directory_rename(self):
        observed = []
        original = os.rename
        def rename(source, destination):
            observed.append({p.name for p in source.iterdir()})
            self.assertFalse(destination.exists())
            return original(source, destination)
        with patch.object(art.os, "rename", rename):
            self.create()
        self.assertEqual(observed, [{"metadata.json", "report.txt"}])

    def test_stale_pending_blocks_writes(self):
        directory, _ = self.store._owner_directory(self.owner, True)
        (directory / (".pending_" + str(uuid.uuid4()))).mkdir()
        with self.assertRaises(art.ArtifactError):
            self.create()

    def test_concurrent_quota_is_not_bypassed(self):
        def attempt(_):
            try:
                return self.store.create(self.owner, "text", "parallel", {"content": "ok"})
            except art.ArtifactError:
                return None
        with patch.object(art, "MAX_OWNER_ARTIFACTS", 1), ThreadPoolExecutor(max_workers=4) as pool:
            results = list(pool.map(attempt, range(8)))
        self.assertEqual(sum(result is not None for result in results), 1)

    def test_invalid_payload_and_no_arbitrary_python_objects(self):
        cyclic = {}
        cyclic["self"] = cyclic
        for kind, data in (("text", {"content": 1}), ("text", {"content": "x", "extra": True}),
                           ("json", {"data": float("nan")}), ("json", {"data": object()}),
                           ("json", cyclic), ("csv", {"headers": ["a"], "rows": [[1, 2]]}),
                           ("exe", {}), ("text", {"content": "bad\0data"})):
            with self.subTest(kind=kind, data=type(data)), self.assertRaises(art.ArtifactError):
                self.store.create(self.owner, kind, "invalid", data)

    def test_scripts_are_bytes_never_executed_and_download_only(self):
        record = self.create("python", "script", {"content": "raise RuntimeError('must never run')\n"})
        path, meta = self.resolve(record)
        self.assertTrue(path.read_bytes().startswith(b"raise RuntimeError"))
        headers = meta["download_headers"]
        self.assertEqual(headers["Content-Type"], "application/octet-stream")
        self.assertTrue(headers["Content-Disposition"].startswith("attachment;"))
        self.assertEqual(headers["X-Content-Type-Options"], "nosniff")
        self.assertIn("sandbox", headers["Content-Security-Policy"])

    def test_json_and_csv_roundtrip_and_injection(self):
        record = self.create("json", "data", {"data": {"olá": [1, True, None]}})
        self.assertEqual(json.loads(self.resolve(record)[0].read_bytes()), {"olá": [1, True, None]})
        record = self.create("csv", "table", {"headers": ["Name", "Value"], "rows": [["=1+1", -2], [" \t@SUM(A1)", 3]]})
        data = self.resolve(record)[0].read_text(encoding="utf-8")
        self.assertIn("'=1+1,-2", data)
        self.assertIn("' \t@SUM(A1)", data)
        record = self.store.ingest(self.owner, "provider.csv", b"A,B\n=1+1,+2\n")
        self.assertIn(b"'=1+1,'+2", self.resolve(record)[0].read_bytes())

    def test_invalid_provider_formats(self):
        samples = [("a.exe", b"MZ"), ("a.json", b'{"a":1,"a":2}'), ("a.json", b"NaN"),
                   ("a.json", b"["), ("a.txt", b"\xff"), ("a.txt", b"a\0b"),
                   ("a.csv", b'A,B\n"unfinished'), ("a.xlsx", b"not a zip"),
                   ("a.pptx", b"not a zip"), ("a.png", b"not png"), ("a.pdf", b"not pdf")]
        for filename, raw in samples:
            with self.subTest(filename=filename, raw=raw), self.assertRaises(art.ArtifactError):
                self.store.ingest(self.owner, filename, raw)

    def test_xlsx_reopens_and_text_is_not_formula(self):
        from openpyxl import load_workbook
        payload = {"sheets": [{"name": "Data", "headers": ["=header", "Value"],
            "rows": [["=HYPERLINK(\"https://evil.test\")", 2], ["@name", 3], ["total", None]],
            "formulas": {"B4": "=SUM(B2:B3)"}}]}
        record = self.create("spreadsheet", "numbers", payload)
        path, _ = self.resolve(record)
        workbook = load_workbook(path, data_only=False)
        self.addCleanup(workbook.close)
        sheet = workbook["Data"]
        self.assertEqual(sheet["A1"].data_type, "s")
        self.assertEqual(sheet["A2"].data_type, "s")
        self.assertEqual(sheet["B4"].value, "=SUM(B2:B3)")
        self.assertEqual(sheet.freeze_panes, "A2")
        cached = load_workbook(path, data_only=True)
        self.assertIsNone(cached["Data"]["B4"].value)  # No recalculation claimed.
        cached.close()
        ingested = self.store.ingest(self.owner, "again.xlsx", path.read_bytes())
        self.assertEqual(ingested["sha256"], record["sha256"])

    def test_formula_allowlist_and_syntax(self):
        good = ["=SUM(B2:B3)", "=AVERAGE(B2,B3)", "=MIN(1,MAX(2,3))", "=($B$2+B3)/2", "=-B2+3^2", "=1e-3+2"]
        bad = ['=HYPERLINK("https://evil.test")', '=WEBSERVICE("x")', "=cmd|' /C calc'!A0", "='[external.xlsx]Data'!A1",
               "=Data!A1", "=INDIRECT(A1)", "=SUM()", "=SUM(,A1)", "=1+", "=1 2", "=A1:A2", "=SUM(A2:A1)",
               "=SUM(A1:A2)+A1:A2", "=A0", "=XFD1048576", "=SUM(A:A)", "=SUM(1);2", "=1e999"]
        for formula in good:
            with self.subTest(formula=formula):
                self.assertTrue(art._formula(formula, 4, 2).startswith("="))
        for formula in bad:
            with self.subTest(formula=formula), self.assertRaises(art.ArtifactError):
                art._formula(formula, 4, 2)

    def test_spreadsheet_shape_limits(self):
        samples = [[], [{"name": "Bad/Name", "headers": ["x"], "rows": []}],
                   [{"name": "A", "headers": ["x"], "rows": [[1]], "formulas": {"A1": "=1"}}],
                   [{"name": "A", "headers": ["x"], "rows": [[1]], "formulas": {"A2": "=1", "$A$2": "=2"}}]]
        for sheets in samples:
            with self.subTest(sheets=sheets), self.assertRaises(art.ArtifactError):
                self.create("spreadsheet", "x", {"sheets": sheets})

    def test_pptx_reopens_with_styled_title_bullets(self):
        from pptx import Presentation
        record = self.create("slides", "briefing", {"slides": [{"title": "Local title", "bullets": ["First item", "Second item"]}]})
        path, _ = self.resolve(record)
        deck = Presentation(path)
        self.assertEqual(len(deck.slides), 1)
        self.assertEqual(deck.slides[0].shapes.title.text, "Local title")
        self.assertIn("Second item", deck.slides[0].placeholders[1].text)
        self.assertEqual(deck.slides[0].shapes.title.text_frame.paragraphs[0].font.size.pt, 34)
        with zipfile.ZipFile(path) as package:
            self.assertFalse(any("printerSettings" in n or "thumbnail" in n for n in package.namelist()))
        self.store.ingest(self.owner, "again.pptx", path.read_bytes())

    def test_office_zip_path_macros_entities_and_external_links(self):
        malicious = [rewrite_zip(self.xlsx_bytes, extra={"../escape": b"bad"}),
                     rewrite_zip(self.xlsx_bytes, extra={"xl/vbaProject.bin": b"macro"}),
                     rewrite_zip(self.xlsx_bytes, extra={"xl/externalLinks/externalLink1.xml": b"<x/>"}),
                     rewrite_zip(self.xlsx_bytes, {"xl/workbook.xml": lambda raw: b'<!DOCTYPE x [<!ENTITY x "boom">]>' + raw}),
                     rewrite_zip(self.xlsx_bytes, {"xl/_rels/workbook.xml.rels": lambda raw: raw.replace(b'Target="', b'TargetMode="External" Target="', 1)}),
                     rewrite_zip(self.xlsx_bytes, {"xl/worksheets/sheet1.xml": lambda raw: raw.replace(b"SUM(B2:B3)", b"WEBSERVICE(1)")})]
        for raw in malicious:
            with self.subTest(size=len(raw)), self.assertRaises(art.ArtifactError):
                self.store.ingest(self.owner, "bad.xlsx", raw)

    def test_zip_size_symlink_and_duplicate(self):
        with patch.object(art, "MAX_ZIP_EXPANDED_BYTES", 100), self.assertRaises(art.ArtifactError):
            self.store.ingest(self.owner, "bad.xlsx", self.xlsx_bytes)
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            duplicate = rewrite_zip(self.xlsx_bytes, extra={"xl/workbook.xml": b"<x/>"})
        with self.assertRaises(art.ArtifactError):
            self.store.ingest(self.owner, "bad.xlsx", duplicate)
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as package:
            info = zipfile.ZipInfo("xl/workbook.xml")
            info.create_system = 3
            info.external_attr = (stat.S_IFLNK | 0o777) << 16
            package.writestr(info, b"/outside")
        with self.assertRaises(art.ArtifactError):
            self.store.ingest(self.owner, "bad.xlsx", buffer.getvalue())

    def test_website_escapes_and_downloads(self):
        record = self.create("website", "site", {"title": "<script>alert(1)</script>", "sections": [{"heading": "Intro", "text": "A & B", "bullets": ["One"]}]})
        path, meta = self.resolve(record)
        source = path.read_text(encoding="utf-8")
        self.assertIn("&lt;script&gt;", source)
        self.assertNotIn("<script>", source)
        self.assertEqual(meta["mime"], "text/html")
        self.store.ingest(self.owner, "site.html", path.read_bytes())

    def test_website_rejects_active_remote_and_css_escapes(self):
        samples = ['<script>alert(1)</script>', '<img src="https://evil.test/a">', '<p onclick="evil()">x</p>',
                   '<iframe srcdoc="x"></iframe>', '<style>body{background:url(https://evil.test)}</style>',
                   '<style>@import "https://evil.test";</style>', '<style>p{background:u\\72l(x)}</style>',
                   '<meta http-equiv="refresh" content="0;url=https://evil.test">', '<object data="file:///x"></object>',
                   '<p><b>unclosed</p>', '<p style="color:red">x</p>']
        for source in samples:
            with self.subTest(source=source), self.assertRaises(art.ArtifactError):
                self.store.ingest(self.owner, "site.html", source.encode())

    def test_png_and_pdf_local_ingestion(self):
        from PIL import Image
        from pypdf import PdfWriter
        png = io.BytesIO()
        Image.new("RGB", (8, 8), (20, 80, 90)).save(png, format="PNG")
        record = self.store.ingest(self.owner, "picture.png", png.getvalue())
        self.assertEqual(self.resolve(record)[1]["mime"], "image/png")
        pdf = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=400, height=400)
        writer.write(pdf)
        record = self.store.ingest(self.owner, "document.pdf", pdf.getvalue())
        self.assertEqual(self.resolve(record)[1]["mime"], "application/pdf")

    def test_active_pdf_and_png_trailing_content_rejected(self):
        from PIL import Image
        from pypdf import PdfWriter
        pdf = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=400, height=400)
        writer.add_js("app.alert('bad')")
        writer.write(pdf)
        with self.assertRaises(art.ArtifactError):
            self.store.ingest(self.owner, "active.pdf", pdf.getvalue())
        png = io.BytesIO()
        Image.new("RGB", (2, 2)).save(png, format="PNG")
        with self.assertRaises(art.ArtifactError):
            self.store.ingest(self.owner, "picture.png", png.getvalue() + b"<script>bad</script>")

    def test_ticket_lifetime_hash_only_and_restart(self):
        with patch.object(art.time, "time", return_value=1_800_000_000):
            record = self.create()
        token = record["mobile_url"].rsplit("/", 1)[1]
        self.assertEqual(len(token), 43)
        ticket_path = self.store.tickets / (record["id"] + ".json")
        ticket = json.loads(ticket_path.read_bytes())
        self.assertNotIn(token.encode(), ticket_path.read_bytes())
        self.assertEqual(ticket["token_sha256"], hashlib.sha256(token.encode()).hexdigest())
        self.assertEqual(ticket["expires_at"], 1_800_000_900)
        for path in self.store.storage.rglob("*"):
            if path.is_file():
                self.assertNotIn(token.encode(), path.read_bytes())
        reloaded = art.ArtifactStore(self.root)
        with patch.object(art.time, "time", return_value=1_800_000_899):
            path, meta = reloaded.resolve_ticket(record["id"], token)
            self.assertEqual(path.read_bytes(), b"hello")
            self.assertEqual(meta["expires_at"], record["expires_at"])
        with patch.object(art.time, "time", return_value=1_800_000_900), self.assertRaises(art.ArtifactError):
            reloaded.resolve_ticket(record["id"], token)

    def test_wrong_token_never_opens_owner_artifact(self):
        record = self.create()
        with patch.object(self.store, "_record", side_effect=AssertionError("Must not access owner data")):
            with self.assertRaises(art.ArtifactError):
                self.store.resolve_ticket(record["id"], "A" * 43)
        for token in ("../x", "", None, "a" * 100):
            with self.subTest(token=token), self.assertRaises(art.ArtifactError):
                self.store.resolve_ticket(record["id"], token)

    def test_ticket_cross_owner_issue_reissue_and_cross_artifact(self):
        record = self.create()
        token = record["mobile_url"].rsplit("/", 1)[1]
        with self.assertRaises(art.ArtifactError):
            self.store.issue_ticket("intruder", record["id"])
        other = self.create(name="second")
        with self.assertRaises(art.ArtifactError):
            self.store.resolve_ticket(other["id"], token)
        new_url = self.store.issue_ticket(self.owner, record["id"])
        self.assertNotEqual(new_url, record["mobile_url"])
        with self.assertRaises(art.ArtifactError):
            self.store.resolve_ticket(record["id"], token)
        with patch.object(art.hmac, "compare_digest", wraps=art.hmac.compare_digest) as compare:
            self.store.resolve_ticket(record["id"], new_url.rsplit("/", 1)[1])
            self.assertTrue(compare.called)
        with self.assertRaises(art.ArtifactError):
            self.store.issue_ticket(self.owner, record["id"], ttl=3600)

    def test_ticket_metadata_path_and_expiry_tampering(self):
        record = self.create()
        token = record["mobile_url"].rsplit("/", 1)[1]
        ticket_path = self.store.tickets / (record["id"] + ".json")
        original = ticket_path.read_bytes()
        for key, value in (("owner_hash", "../../escape"), ("id", str(uuid.uuid4())), ("expires_at", 9999999999), ("schema_version", True)):
            content = json.loads(original)
            content[key] = value
            ticket_path.write_text(json.dumps(content), encoding="utf-8")
            with self.subTest(key=key), self.assertRaises(art.ArtifactError):
                self.store.resolve_ticket(record["id"], token)
            ticket_path.write_bytes(original)


if __name__ == "__main__":
    unittest.main(verbosity=2)
